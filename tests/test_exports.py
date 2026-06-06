import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from skp_li_saves.export_html import export_html
from skp_li_saves.export_xlsx import export_xlsx


REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "src"))


class ExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.records = [
            {
                "post_id": "urn:li:activity:1",
                "post_url": "https://www.linkedin.com/feed/update/urn:li:activity:1/",
                "author_name": "Ada Lovelace",
                "author_role": "Founder",
                "published_at": "2026-06-01T12:00:00Z",
                "text_clean": "Hello world from LinkedIn",
                "text_raw": "Hello world from LinkedIn\nLike\nComment\nShare\n",
                "summary": "A short intro post",
                "theme_primary": "AI",
                "theme_secondary": "Automation",
                "source_type": "linkedin_export",
                "source_run_id": "run-123",
                "schema_version": "1",
            },
            {
                "post_id": "urn:li:activity:2",
                "post_url": "https://www.linkedin.com/feed/update/urn:li:activity:2/",
                "author_name": "Grace Hopper",
                "author_role": "Researcher",
                "published_at": "2026-06-02T08:30:00Z",
                "text_clean": "Second post with a longer body for excerpt rendering",
                "text_raw": "Second post with a longer body for excerpt rendering",
                "summary": "Another saved post",
                "theme_primary": "AI",
                "theme_secondary": "",
                "source_type": "linkedin_export",
                "source_run_id": "run-123",
                "schema_version": "1",
            },
        ]

    def test_html_export_writes_self_contained_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "library.html"
            export_html(self.records, output_path, title="Saved Posts Library")
            html = output_path.read_text(encoding="utf-8")

            self.assertIn("Saved Posts Library", html)
            self.assertIn("saved-posts-data", html)
            self.assertIn("localStorage", html)
            self.assertIn("theme-filters", html)
            self.assertIn("Ada Lovelace", html)
            self.assertIn("Open on LinkedIn", html)
            self.assertIn("Search author, role, summary, or text", html)
            self.assertIn("Sort: recency", html)

    def test_xlsx_export_writes_expected_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "library.xlsx"
            export_xlsx(self.records, output_path)

            workbook = load_workbook(output_path)
            self.assertEqual(workbook.sheetnames, ["Saved Posts", "By Theme"])

            posts = workbook["Saved Posts"]
            self.assertEqual(posts.freeze_panes, "A2")
            self.assertEqual(posts.auto_filter.ref, posts.dimensions)
            self.assertEqual([cell.value for cell in posts[1]], [
                "post_id",
                "author_name",
                "author_role",
                "published_at",
                "summary",
                "theme_primary",
                "theme_secondary",
                "text_clean",
                "post_url",
            ])
            self.assertEqual(posts[2][1].value, "Ada Lovelace")
            self.assertEqual(posts[2][8].hyperlink.target, "https://www.linkedin.com/feed/update/urn:li:activity:1/")

            themes = workbook["By Theme"]
            self.assertEqual(themes.freeze_panes, "A2")
            self.assertEqual([cell.value for cell in themes[1]], ["theme", "count"])
            rows = list(themes.iter_rows(min_row=2, values_only=True))
            self.assertEqual(rows, [("AI", 2), ("Automation", 1)])

    def test_cli_exports_html_from_json_array_input(self) -> None:
        env = os.environ.copy()
        env["PYTHONPATH"] = str(REPO_ROOT / "src")
        with tempfile.TemporaryDirectory() as tmpdir:
            tmpdir_path = Path(tmpdir)
            input_path = tmpdir_path / "normalized.json"
            input_path.write_text(json.dumps(self.records, ensure_ascii=False), encoding="utf-8")

            result = subprocess.run(
                [sys.executable, "-m", "skp_li_saves", str(input_path), "--format", "html"],
                cwd=REPO_ROOT,
                env=env,
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(result.returncode, 0, msg=result.stderr)
            html_path = tmpdir_path / "normalized.html"
            self.assertTrue(html_path.exists())
            html = html_path.read_text(encoding="utf-8")
            self.assertIn("Ada Lovelace", html)
            self.assertIn("Saved Posts", html)

    def test_cli_exports_xlsx_from_jsonl_input(self) -> None:
        env = os.environ.copy()
        env["PYTHONPATH"] = str(REPO_ROOT / "src")
        with tempfile.TemporaryDirectory() as tmpdir:
            tmpdir_path = Path(tmpdir)
            input_path = tmpdir_path / "normalized.jsonl"
            input_path.write_text("\n".join(json.dumps(record, ensure_ascii=False) for record in self.records) + "\n", encoding="utf-8")
            output_path = tmpdir_path / "library.xlsx"

            result = subprocess.run(
                [sys.executable, "-m", "skp_li_saves", str(input_path), "--format", "xlsx", "--output", str(output_path)],
                cwd=REPO_ROOT,
                env=env,
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(result.returncode, 0, msg=result.stderr)
            self.assertTrue(output_path.exists())
            workbook = load_workbook(output_path)
            self.assertEqual(workbook["By Theme"]["A2"].value, "AI")


if __name__ == "__main__":
    unittest.main()
