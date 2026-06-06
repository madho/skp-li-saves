"""XLSX export for normalized saved posts."""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from .export_common import record_to_export_dict, theme_counts


SAVED_POST_HEADERS = [
    "post_id",
    "author_name",
    "author_role",
    "published_at",
    "summary",
    "theme_primary",
    "theme_secondary",
    "text_clean",
    "post_url",
]


THEME_HEADERS = ["theme", "count"]


def _apply_header_row(sheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True)


def export_xlsx(records: Iterable[dict[str, Any]], output_path: Path) -> Path:
    """Write an XLSX workbook with saved posts and theme counts."""

    normalized = [record_to_export_dict(record) for record in records]

    workbook = Workbook()
    posts_sheet = workbook.active
    posts_sheet.title = "Saved Posts"
    posts_sheet.freeze_panes = "A2"
    posts_sheet.append(SAVED_POST_HEADERS)

    for record in normalized:
        row = [record.get(column, "") for column in SAVED_POST_HEADERS]
        posts_sheet.append(row)
        url = record.get("post_url", "")
        if url:
            url_cell = posts_sheet.cell(row=posts_sheet.max_row, column=len(SAVED_POST_HEADERS))
            url_cell.hyperlink = url
            url_cell.style = "Hyperlink"

    posts_sheet.auto_filter.ref = posts_sheet.dimensions
    _apply_header_row(posts_sheet)

    widths = {
        "A": 24,
        "B": 24,
        "C": 24,
        "D": 24,
        "E": 32,
        "F": 20,
        "G": 20,
        "H": 50,
        "I": 48,
    }
    for column, width in widths.items():
        posts_sheet.column_dimensions[column].width = width

    theme_sheet = workbook.create_sheet("By Theme")
    theme_sheet.freeze_panes = "A2"
    theme_sheet.append(THEME_HEADERS)
    for theme, count in theme_counts(normalized):
        theme_sheet.append([theme, count])
    theme_sheet.auto_filter.ref = theme_sheet.dimensions
    _apply_header_row(theme_sheet)
    theme_sheet.column_dimensions["A"].width = 28
    theme_sheet.column_dimensions["B"].width = 12

    workbook.save(output_path)
    return output_path
